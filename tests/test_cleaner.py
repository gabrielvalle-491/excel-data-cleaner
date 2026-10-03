from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from data_cleaner.generate_messy_data import generate
from data_cleaner.pipeline import clean_dataframe, load_config, read_table
from data_cleaner.report import write_report
from data_cleaner.rules import clean_amount, clean_country, clean_date, clean_email, clean_name, clean_phone

CONFIG = load_config(Path(__file__).parent.parent / "config.yaml")


@pytest.mark.parametrize("raw, expected", [
    ("  juan   PÉREZ ", "Juan Pérez"),
    ("maría DE LA fuente", "María de la Fuente"),
])
def test_clean_name(raw, expected):
    assert clean_name(raw) == (expected, None)


def test_clean_email_fixes_case_spaces_and_typos():
    assert clean_email(" Ana.Lopez@GMAIL.COM ") == ("ana.lopez@gmail.com", None)
    value, issue = clean_email("ana@gmial.com")
    assert value == "ana@gmail.com" and "typo" in issue
    assert clean_email("not-an-email")[1] == "Invalid email"


@pytest.mark.parametrize("raw", ["+54 9 2657 351236", "0265715351236", "(2657) 351236", "2657351236",
                                 "5492657351236"])
def test_clean_phone_argentina_variants(raw):
    assert clean_phone(raw, "Argentina") == ("+5492657351236", None)


def test_clean_phone_other_country():
    assert clean_phone("55 1234 5678", "Mexico") == ("+525512345678", None)


@pytest.mark.parametrize("raw", ["2025-03-07", "07/03/2025", "07-03-2025", "Mar 07, 2025", "07.03.2025"])
def test_clean_date_formats(raw):
    assert clean_date(raw) == ("2025-03-07", None)


def test_clean_date_rejects_impossible_dates():
    assert clean_date("31/02/2025")[1] == "Unrecognized date"


@pytest.mark.parametrize("raw, expected", [("$ 1.234,50", 1234.5), ("1,234.50", 1234.5), ("USD 99.9", 99.9)])
def test_clean_amount(raw, expected):
    assert clean_amount(raw) == (expected, None)


def test_clean_country_aliases():
    assert clean_country("ARG") == ("Argentina", None)
    assert clean_country("méxico") == ("Mexico", None)
    assert clean_country("Atlantis")[1] == "Unknown country"


def test_pipeline_removes_duplicates_and_reports_issues():
    df = pd.DataFrame({
        "Nombre": ["ana lópez", "ANA LÓPEZ", "Pedro Gil", None],
        "E-mail": ["ana@gmail.com", " ANA@GMAIL.COM", "bad-email", None],
        "Telefono": ["2657351236", "+54 9 2657 351236", "11 5555 4444", None],
        "País": ["Argentina", "arg", "AR", None],
        "Fecha alta": ["2025-01-01", "01/01/2025", "nope", None],
        "Monto compra": ["100", "100", "50", None],
        "Ciudad": ["Villa Mercedes", "villa  mercedes", "-", None],
    })
    result = clean_dataframe(df, CONFIG)
    assert result.stats["rows_in"] == 3  # the fully blank row is dropped
    assert result.stats["duplicates_removed"] == 1
    assert list(result.cleaned["email"]) == ["ana@gmail.com", "bad-email"]
    issues = set(result.issues["issue"])
    assert {"Invalid email", "Unrecognized date"} <= issues


def test_end_to_end_report(tmp_path):
    source = generate(tmp_path / "messy.xlsx", rows=200)
    result = clean_dataframe(read_table(source), CONFIG)
    assert result.stats["duplicates_removed"] > 0
    assert result.cleaned["email"].dropna().str.match(r"^[a-z0-9._%+\-]+@").all()
    assert result.cleaned["phone"].dropna().str.startswith("+").all()
    path = write_report(result, tmp_path / "report.xlsx")
    assert load_workbook(path).sheetnames == ["Summary", "Clean data", "Issues", "Duplicates"]


def test_clean_phone_flags_country_mismatch():
    value, issue = clean_phone("542657277863", "Chile")
    assert value == "+5492657277863"
    assert "does not match country" in issue
