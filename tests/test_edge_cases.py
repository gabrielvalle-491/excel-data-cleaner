import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from data_cleaner.cli import main
from data_cleaner.pipeline import clean_dataframe, load_config
from data_cleaner.report import write_report
from data_cleaner.rules import clean_amount, clean_date, clean_phone

ROOT = Path(__file__).parent.parent
CONFIG = load_config(ROOT / "config.yaml")
HEADERS = ["Nombre", "E-mail", "Telefono", "País", "Fecha alta", "Monto compra", "Ciudad"]


def test_help_exits_zero_and_shows_module_name():
    proc = subprocess.run([sys.executable, "-m", "data_cleaner", "--help"], cwd=ROOT,
                          capture_output=True, text=True)
    assert proc.returncode == 0
    assert proc.stdout.startswith("usage: python -m data_cleaner")


def test_missing_input_or_config_returns_exit_code_1(tmp_path, capsys):
    assert main([str(tmp_path / "missing.xlsx")]) == 1
    source = tmp_path / "in.csv"
    source.write_text("Nombre,E-mail\nAna,ana@gmail.com\n", encoding="utf-8")
    assert main([str(source), "-c", str(tmp_path / "missing.yaml")]) == 1
    assert "File not found" in capsys.readouterr().err


def test_headers_only_file_produces_empty_report(tmp_path):
    result = clean_dataframe(pd.DataFrame(columns=HEADERS), CONFIG)
    assert result.stats == {"rows_in": 0, "rows_out": 0, "duplicates_removed": 0,
                            "issues": 0, "rows_with_issues": 0}
    path = write_report(result, tmp_path / "empty.xlsx")
    assert load_workbook(path).sheetnames == ["Summary", "Clean data", "Issues", "Duplicates"]


def test_missing_columns_are_reported_without_a_row_number():
    result = clean_dataframe(pd.DataFrame({"Nombre": ["Ana López"]}), CONFIG)
    missing = result.issues[result.issues["issue"] == "Column missing in file"]
    assert set(missing["column"]) == {"email", "phone", "country", "signup_date", "purchase_amount", "city"}
    assert (missing["row"] == "-").all()
    assert result.stats["rows_with_issues"] == 0


@pytest.mark.parametrize("rule, raw, expected", [
    (clean_amount, "abc", ("abc", "Invalid amount")),
    (clean_amount, "$ -1.500,00", (-1500.0, "Negative amount")),
    (clean_amount, "n/a", (None, "Missing amount")),
    (clean_date, "01/01/1800", ("01/01/1800", "Date out of range")),
    (clean_phone, "123", ("+54123", "Phone has an invalid length")),
])
def test_malformed_and_boundary_values(rule, raw, expected):
    assert rule(raw) == expected
