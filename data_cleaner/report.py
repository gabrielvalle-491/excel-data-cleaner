"""Write the cleaning result as a formatted Excel report."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from data_cleaner.pipeline import CleanResult

HEADER_FILL = PatternFill("solid", fgColor="375623")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_report(result: CleanResult, output: str | Path) -> Path:
    """Write the Summary, Clean data, Issues and Duplicates sheets and return the report path."""
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)

    issue_counts = (result.issues.groupby(["column", "issue"]).size().reset_index(name="count")
                    .sort_values("count", ascending=False))
    summary = pd.DataFrame({
        "Metric": ["Rows received", "Rows after cleaning", "Duplicates removed", "Issues found",
                   "Rows with at least one issue"],
        "Value": [result.stats["rows_in"], result.stats["rows_out"], result.stats["duplicates_removed"],
                  result.stats["issues"], result.stats["rows_with_issues"]],
    })

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        issue_counts.to_excel(writer, sheet_name="Summary", index=False, startrow=len(summary) + 2)
        result.cleaned.to_excel(writer, sheet_name="Clean data", index=False)
        result.issues.to_excel(writer, sheet_name="Issues", index=False)
        result.duplicates.to_excel(writer, sheet_name="Duplicates", index=False)

    wb = load_workbook(output)
    for ws in wb.worksheets:
        header_rows = [1] + ([len(summary) + 3] if ws.title == "Summary" else [])
        for r in header_rows:
            for cell in ws[r]:
                if cell.value is not None:
                    cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for idx, column in enumerate(ws.columns, start=1):
            width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(width + 3, 50)
        if ws.title != "Summary":
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for column in ws.iter_cols(min_row=2):
            if any(isinstance(c.value, float) for c in column):
                for cell in column:
                    cell.number_format = "#,##0.00"
    wb.save(output)
    return output
